#!/usr/bin/env python3
"""Round-trip a newly added hotel and per-night reservations through Excel."""
import importlib.util
import json
import os
import shutil
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook

from lib_py import Suite, finish

REPO = Path(__file__).resolve().parent.parent
TMP = Path("/tmp/wt/pytest")
SRC = REPO / "Wedding_Expense_Tracker_Dec2026.xlsx"
TMP.mkdir(parents=True, exist_ok=True)

spec = importlib.util.spec_from_file_location(
    "apply_payload", REPO / "tools" / "apply_payload.py")
apply = importlib.util.module_from_spec(spec)
spec.loader.exec_module(apply)

workbook = TMP / "dynamic-hotel.xlsx"
export_path = TMP / "dynamic-hotel-export.json"
shutil.copy2(SRC, workbook)
start = date(2026, 12, 8)
nights = [(start + timedelta(days=i)).strftime("%d-%b-%Y") for i in range(8)]
needed = [1, 2, 1, 0, 2, 2, 1, 1]
payload = {
    "v": 2,
    "mode": "full",
    "author": "Dynamic hotel test",
    "rooms": [
        {
            "hotel": "Nirmal",
            "nights": ["08-Dec-2026", "09-Dec-2026", "10-Dec-2026",
                       "11-Dec-2026", "12-Dec-2026", "13-Dec-2026"],
            "needed": [1, 2, 3, 4, 5, 6],
        },
        {
            "hotel": "Palm Grove",
            "id": "palm-grove",
            "totalRooms": 2,
            "roomsBooked": 2,
            "roomNos": ["PG-1", "PG-2"],
            "checkIn": nights[0],
            "checkOut": (start + timedelta(days=len(nights))).strftime("%d-%b-%Y"),
            "checkoutTime": "",
            "nights": nights,
            "needed": needed,
        },
    ],
    "guests": [{
        "hotel": "Palm Grove", "room": 2, "night": nights[-1],
        "slot": 3, "name": "Round Trip Guest",
    }],
}
result = apply.apply_payloads([payload], workbook, make_backup=False)
proc = subprocess.run(
    [sys.executable, str(REPO / "tools" / "export_site_data.py"),
     "--workbook", str(workbook), "--out", str(export_path)],
    capture_output=True, text=True,
    env={**os.environ, "PYTHONPATH": os.pathsep.join(sys.path)},
)
S = Suite("INTEGRATION / dynamic hotel Excel backup")
S.eq(result["skipped"], [], "hotel and guest payload applied without skips")
S.eq(result["guests"], 1, "new hotel guest was written")
S.eq(proc.returncode, 0, "workbook export completed", proc.stderr)
if proc.returncode == 0:
    exported = json.loads(export_path.read_text())
    hotel = next(h for h in exported["hotels"] if h["id"] == "palm-grove")
    S.eq(hotel["nights"], nights, "export supports more than six dates")
    S.eq(hotel["needed"], needed, "nightly pre-booked counts round-trip")
    S.eq(hotel["totalRooms"], 2, "room capacity round-trips")
    S.eq(hotel["roomNos"], ["PG-1", "PG-2"], "room numbers round-trip")
    S.eq(hotel["grid"][1][-1][2], "Round Trip Guest",
         "guest details round-trip on the last night and third slot")
    existing = next(h for h in exported["hotels"] if h["id"] == "nirmal")
    S.eq(existing["needed"], [1, 2, 3, 4, 5, 6],
         "existing hotel reservation counts update independently by date")

sheet = load_workbook(workbook)["Palm Grove"]
S.ok("Guest details complete" in str(sheet["J39"].value)
     or any("Guest details complete" in str(cell.value)
            for row in sheet.iter_rows() for cell in row),
     "the created hotel sheet has the guest-detail status formula")
finish([S])
